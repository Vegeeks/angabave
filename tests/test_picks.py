"""Pruebas de la lectura del Excel semanal."""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pytest

from quiniela.picks import (
    ErrorFormatoExcel,
    ErrorPicks,
    Hoja,
    archivos_por_semana,
    clave_participante,
    escribir_hoja,
    leer_hoja,
    leer_picks,
    numero_semana,
    semana_en_el_nombre,
)

from .conftest import ENFRENTAMIENTOS_S2, PICKS_S2, crear_excel


def test_lee_los_16_enfrentamientos_en_orden(excel_s2: Path):
    partidos, _ = leer_picks(excel_s2)
    assert len(partidos) == 16
    assert [(p.visitante, p.local) for p in partidos] == ENFRENTAMIENTOS_S2
    assert partidos[0].clave == "Lions@Bills"
    # Del Excel no salen marcadores: son partidos sin resultado.
    assert all(not p.finalizado and p.ganador is None for p in partidos)


def test_lee_los_picks_de_cada_participante(excel_s2: Path):
    _, picks = leer_picks(excel_s2)
    assert list(picks) == list(PICKS_S2)
    assert all(len(elegidos) == 16 for elegidos in picks.values())
    assert picks["Ismael Reyna"][9] == "Jaguars"
    assert picks["Angel D Luffy"][9] == "Broncos"


def test_ignora_la_columna_de_aciertos_totales(excel_s2: Path, caplog):
    with caplog.at_level(logging.INFO, logger="quiniela.picks"):
        partidos, picks = leer_picks(excel_s2)
    assert len(partidos) == 16
    assert all("99" not in elegidos for elegidos in picks.values())
    assert "Aciertos Totales" in caplog.text


def test_el_encabezado_acepta_commanders_y_devuelve_washington(tmp_path: Path):
    enfrentamientos = [("COMMANDERS" if v == "Washington" else v, l) for v, l in ENFRENTAMIENTOS_S2]
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", enfrentamientos=enfrentamientos)
    partidos, _ = leer_picks(ruta)
    assert partidos[12].visitante == "Washington"


def test_falta_un_pick(tmp_path: Path):
    picks = {nombre: list(elegidos) for nombre, elegidos in PICKS_S2.items()}
    picks["Angel D Luffy"][7] = ""  # se queda con 15
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", picks=picks)

    with pytest.raises(ErrorPicks) as excepcion:
        leer_picks(ruta)
    mensaje = str(excepcion.value)
    assert "Angel D Luffy" in mensaje
    assert "15 picks de 16" in mensaje
    assert "Browns@Buccaneers" in mensaje


def test_sobran_picks_porque_hay_una_columna_de_mas(tmp_path: Path):
    picks = {nombre: list(elegidos) for nombre, elegidos in PICKS_S2.items()}
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", picks=picks)
    libro = openpyxl.load_workbook(ruta)
    hoja = libro.active
    hoja.cell(row=3, column=18, value="Rams")  # invade la columna de totales
    libro.save(ruta)

    # La columna de totales se ignora: el pick de más no se cuela en la tabla.
    _, leidos = leer_picks(ruta)
    assert len(leidos["Ismael Reyna"]) == 16


def test_pick_que_no_corresponde_a_ese_partido(tmp_path: Path):
    picks = {nombre: list(elegidos) for nombre, elegidos in PICKS_S2.items()}
    picks["Ismael Reyna"][0] = "Cowboys"  # el partido es Lions@Bills
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", picks=picks)

    with pytest.raises(ErrorPicks) as excepcion:
        leer_picks(ruta)
    assert "eligió Cowboys en Lions@Bills" in str(excepcion.value)


def test_pick_ilegible(tmp_path: Path):
    picks = {nombre: list(elegidos) for nombre, elegidos in PICKS_S2.items()}
    picks["Ismael Reyna"][0] = "Bilis"
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", picks=picks)

    with pytest.raises(ErrorPicks, match="Bilis"):
        leer_picks(ruta)


def test_participante_duplicado(tmp_path: Path):
    picks = dict(PICKS_S2)
    picks["Ismael  Reyna "] = list(PICKS_S2["Ismael Reyna"])
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", picks=picks)

    with pytest.raises(ErrorPicks, match="duplicado"):
        leer_picks(ruta)


def test_acentos_inconsistentes_son_el_mismo_participante():
    assert clave_participante("Chilangos Norteños") == clave_participante("Chilangos Nortenos")
    assert clave_participante("  Angél   D  Luffy ") == clave_participante("Angel D Luffy")


def test_avisa_en_el_log_los_espacios_que_corrigio(tmp_path: Path, caplog):
    picks = {"  Sol   Vega ": list(PICKS_S2["Ismael Reyna"])}
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", picks=picks)

    with caplog.at_level(logging.INFO, logger="quiniela.picks"):
        _, leidos = leer_picks(ruta)
    assert list(leidos) == ["Sol Vega"]
    assert "corregí espacios o acentos" in caplog.text


def test_columnas_separadoras_vacias_se_ignoran(tmp_path: Path):
    ruta = tmp_path / "Semana_02.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.cell(row=2, column=1, value="Semana 2")
    # Un partido por cada dos columnas, con una columna vacía de por medio.
    for indice, (visitante, local) in enumerate(ENFRENTAMIENTOS_S2[:3]):
        columna = 2 + indice * 2
        hoja.cell(row=1, column=columna, value=visitante)
        hoja.cell(row=2, column=columna, value=local)
    hoja.cell(row=2, column=8, value="Aciertos Totales")
    hoja.cell(row=3, column=1, value="Sol Vega")
    for indice, equipo in enumerate(["Bills", "Falcons", "Ravens"]):
        hoja.cell(row=3, column=2 + indice * 2, value=equipo)
    libro.save(ruta)

    partidos, picks = leer_picks(ruta)
    assert len(partidos) == 3
    assert picks["Sol Vega"] == ["Bills", "Falcons", "Ravens"]


def test_la_semana_del_nombre_debe_coincidir_con_la_del_contenido(tmp_path: Path):
    ruta = crear_excel(tmp_path / "Semana_03.xlsx", semana=2)
    with pytest.raises(ErrorPicks, match="Semana 2"):
        leer_picks(ruta)


def test_numero_de_semana_desde_el_nombre(tmp_path: Path):
    assert numero_semana(Path("Semana_02.xlsx")) == 2
    assert numero_semana(Path("Semana 11.xlsx")) == 11
    with pytest.raises(ErrorPicks, match="número de semana"):
        numero_semana(Path("picks.xlsx"))


def test_archivo_sin_enfrentamientos(tmp_path: Path):
    ruta = tmp_path / "Semana_02.xlsx"
    libro = openpyxl.Workbook()
    libro.active.cell(row=1, column=1, value="Semana 2")
    libro.save(ruta)
    with pytest.raises(ErrorFormatoExcel, match="enfrentamientos"):
        leer_picks(ruta)


def test_equipo_repetido_en_la_semana(tmp_path: Path):
    enfrentamientos = list(ENFRENTAMIENTOS_S2)
    enfrentamientos[1] = ("Lions", "Falcons")  # Lions ya juega en la columna anterior
    ruta = crear_excel(tmp_path / "Semana_02.xlsx", enfrentamientos=enfrentamientos)
    with pytest.raises(ErrorFormatoExcel, match="más de un partido"):
        leer_picks(ruta)


def test_archivo_inexistente(tmp_path: Path):
    with pytest.raises(ErrorPicks, match="No existe"):
        leer_picks(tmp_path / "no_esta.xlsx")


# --- la hoja canónica y los nombres ------------------------------------------------


@pytest.mark.parametrize("nombre, semana", [
    ("Quiniela 3.xlsx", 3),
    ("Quiniela.3.Jueves.pdf", 3),
    ("Semana_03.xlsx", 3),
    ("semana 12 final.pdf", 12),
    ("Quiniela 2026.xlsx", None),
    ("picks.xlsx", None),
])
def test_semana_en_el_nombre(nombre, semana):
    assert semana_en_el_nombre(nombre) == semana


def test_la_hoja_que_se_escribe_se_lee_identica(tmp_path: Path):
    original = leer_hoja(crear_excel(tmp_path / "Semana_02.xlsx"))
    copia = escribir_hoja(tmp_path / "Semana_02_copia.xlsx", 2, original.partidos, original.picks)
    releida = leer_hoja(copia)
    assert releida == Hoja(2, original.partidos, original.picks)


def test_leer_hoja_dice_la_semana(excel_s2: Path):
    assert leer_hoja(excel_s2).semana == 2


@pytest.mark.parametrize("nombre", ["<b>Ana</b>", "=1+1", "@Ana", "Ana\u200bLópez", "A" * 41])
def test_el_lector_rechaza_lo_que_no_es_un_nombre(tmp_path: Path, nombre: str):
    picks = {**PICKS_S2, nombre: list(PICKS_S2["Ismael Reyna"])}
    with pytest.raises(ErrorPicks, match="nombre"):
        leer_picks(crear_excel(tmp_path / "Semana_02.xlsx", picks=picks))


@pytest.mark.parametrize("nombre", ["Chucho el Destructor", "Chilangos Norteños", "O'Brien",
                                    "Raider79", "Ana-María", "Pérez & Hijos", "J. López (2)"])
def test_el_lector_acepta_nombres_de_verdad(tmp_path: Path, nombre: str):
    picks = {**PICKS_S2, nombre: list(PICKS_S2["Ismael Reyna"])}
    _, leidos = leer_picks(crear_excel(tmp_path / "Semana_02.xlsx", picks=picks))
    assert nombre in leidos


def test_archivos_por_semana_prefiere_el_excel_y_ignora_temporales(tmp_path: Path):
    crear_excel(tmp_path / "Semana_02.xlsx")
    crear_excel(tmp_path / "Semana_01.xlsx", semana=1)
    (tmp_path / "Semana_02.pdf").write_bytes(b"%PDF-")
    (tmp_path / ".Semana_03.nuevo.xlsx").write_bytes(b"")
    (tmp_path / "~$Semana_01.xlsx").write_bytes(b"")
    (tmp_path / "notas.txt").write_text("x")
    grupos = archivos_por_semana(tmp_path)
    assert list(grupos) == [1, 2]
    assert [p.name for p in grupos[2]] == ["Semana_02.xlsx", "Semana_02.pdf"]
