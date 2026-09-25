"""Utilidades compartidas por las pruebas.

Aquí se arman los Excel de prueba con el mismo layout que el archivo real:
fila de visitantes, fila de locales debajo, participantes a partir de la
tercera fila y una última columna "Aciertos Totales" con valores basura.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

#: Enfrentamientos reales de la Semana 2 de 2026, en el orden del archivo.
ENFRENTAMIENTOS_S2: list[tuple[str, str]] = [
    ("Lions", "Bills"),
    ("Panthers", "Falcons"),
    ("Saints", "Ravens"),
    ("Vikings", "Bears"),
    ("Bengals", "Texans"),
    ("Steelers", "Patriots"),
    ("Packers", "Jets"),
    ("Browns", "Buccaneers"),
    ("Eagles", "Titans"),
    ("Jaguars", "Broncos"),
    ("Raiders", "Chargers"),
    ("Seahawks", "Cardinals"),
    ("Washington", "Cowboys"),
    ("Dolphins", "49ers"),
    ("Colts", "Chiefs"),
    ("Giants", "Rams"),
]

#: Participantes de prueba. Los nombres son reales, los picks NO: están
#: inventados para cubrir casos concretos. Los de verdad viven en data/picks/.
PICKS_S2: dict[str, list[str]] = {
    "Ismael Reyna": [
        "Bills", "Panthers", "Ravens", "Bears", "Texans", "Patriots", "Packers",
        "Buccaneers", "Eagles", "Jaguars", "Chargers", "Seahawks", "Cowboys",
        "49ers", "Chiefs", "Rams",
    ],
    "Angel D Luffy": [
        "Bills", "Panthers", "Ravens", "Bears", "Texans", "Patriots", "Packers",
        "Buccaneers", "Eagles", "Broncos", "Chargers", "Seahawks", "Cowboys",
        "49ers", "Chiefs", "Rams",
    ],
    "Chilangos Norteños": [
        "Bills", "Panthers", "Ravens", "Bears", "Texans", "Steelers", "Packers",
        "Buccaneers", "Eagles", "Jaguars", "Chargers", "Seahawks", "Cowboys",
        "49ers", "Chiefs", "Rams",
    ],
    "Tristan Mejia": [
        "Lions", "Panthers", "Ravens", "Vikings", "Texans", "Patriots", "Packers",
        "Buccaneers", "Eagles", "Broncos", "Chargers", "Cardinals", "Cowboys",
        "49ers", "Colts", "Rams",
    ],
}


def crear_excel(
    ruta: Path,
    enfrentamientos: list[tuple[str, str]] | None = None,
    picks: dict[str, list[str]] | None = None,
    semana: int = 2,
    con_columna_totales: bool = True,
) -> Path:
    """Escribe un Excel de picks con el layout del archivo real."""
    enfrentamientos = enfrentamientos if enfrentamientos is not None else ENFRENTAMIENTOS_S2
    picks = picks if picks is not None else PICKS_S2

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = f"Semana {semana}"

    for columna, (visitante, local) in enumerate(enfrentamientos, start=2):
        hoja.cell(row=1, column=columna, value=visitante)
        hoja.cell(row=2, column=columna, value=local)
    hoja.cell(row=2, column=1, value=f"Semana {semana}")

    columna_totales = len(enfrentamientos) + 2
    if con_columna_totales:
        hoja.cell(row=2, column=columna_totales, value="Aciertos Totales")

    for fila, (participante, elegidos) in enumerate(picks.items(), start=3):
        hoja.cell(row=fila, column=1, value=participante)
        for columna, pick in enumerate(elegidos, start=2):
            if pick:
                hoja.cell(row=fila, column=columna, value=pick)
        if con_columna_totales:
            # Valor deliberadamente incorrecto: es la columna que hay que ignorar.
            hoja.cell(row=fila, column=columna_totales, value=99)

    libro.save(ruta)
    return ruta


@pytest.fixture
def excel_s2(tmp_path: Path) -> Path:
    """Excel de la Semana 2 completo y bien formado."""
    return crear_excel(tmp_path / "Semana_02.xlsx")


def crear_pdf(
    ruta: Path,
    enfrentamientos=None,
    picks=None,
    semana: int = 2,
    *,
    centrado: bool = False,
) -> Path:
    """Arma un PDF mínimo con la rejilla de la quiniela.

    No es un PDF completo, pero sí tiene lo que lee `quiniela.pdf`: un flujo
    comprimido con fragmentos de texto posicionados. Así las pruebas no
    dependen de un archivo real del organizador.

    Con `centrado` arma la hoja como la manda el organizador cuando trae una
    sola tanda: la rejilla corrida a la derecha y los nombres colgados de una
    misma orilla, de modo que los largos empiezan antes que "Semana N".
    """
    import zlib

    enfrentamientos = enfrentamientos if enfrentamientos is not None else ENFRENTAMIENTOS_S2
    picks = picks if picks is not None else PICKS_S2

    x_nombre = 180.0 if centrado else 20.0
    x_primera = 280.0 if centrado else 120.0
    paso = 30.0

    def x_de(celda: str) -> float:
        """En la hoja centrada los nombres cuelgan de la misma orilla derecha."""
        if not centrado:
            return x_nombre
        return x_nombre + (len("Semana 99") - len(celda)) * 3.0
    y_visitantes, y_locales = 555.0, 517.0

    fragmentos: list[tuple[float, float, str]] = []
    for indice, (visitante, local) in enumerate(enfrentamientos):
        x = x_primera + indice * paso
        fragmentos.append((y_visitantes, x, visitante))
        fragmentos.append((y_locales, x, local))
    etiqueta = f"Semana {semana}"
    fragmentos.append((y_locales, x_de(etiqueta), etiqueta))
    # La columna de totales, a la derecha del último partido.
    x_totales = x_primera + len(enfrentamientos) * paso
    fragmentos.append((y_locales + 8, x_totales, "Aciertos"))
    fragmentos.append((y_locales, x_totales, "Totales"))

    for fila, (participante, elegidos) in enumerate(picks.items()):
        y = y_locales - 14 * (fila + 1)
        fragmentos.append((y, x_de(participante), participante))
        for indice, pick in enumerate(elegidos):
            if pick:
                fragmentos.append((y, x_primera + indice * paso, pick))
        fragmentos.append((y, x_totales, "99"))

    bloques = "".join(
        f"BT /F1 6 Tf 1 0 0 1 {x} {y} Tm [({texto})] TJ ET\n"
        for y, x, texto in fragmentos
    )
    comprimido = zlib.compress(bloques.encode("latin-1"))

    cuerpo = (
        b"%PDF-1.4\n"
        b"1 0 obj <</Type/Catalog/Pages 2 0 R>> endobj\n"
        b"2 0 obj <</Type/Pages/Kids[3 0 R]/Count 1>> endobj\n"
        b"3 0 obj <</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R>> endobj\n"
        b"4 0 obj <</Length " + str(len(comprimido)).encode() + b"/Filter/FlateDecode>>\nstream\n"
        + comprimido
        + b"\nendstream endobj\n"
        b"trailer <</Root 1 0 R>>\n%%EOF\n"
    )
    ruta.write_bytes(cuerpo)
    return ruta
